import os
import sys
import pandas as pd
import argparse

# 确保本地模块优先加载
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from config import INPUT_FILE, OUTPUT_DIR

def clean_data(custom_periods=None, input_file=None, output_prefix=None):
    file_path = input_file or INPUT_FILE
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"输入文件 {file_path} 不存在")
    df = pd.read_excel(file_path)
    
    # 数据清洗
    df = df.dropna(subset=['alarm_first_time'])  # 删除时间缺失记录
    df['GRADE'] = df['GRADE'].str.upper()  # 统一等级格式
    # 将GRADE列重命名为中文
    df = df.rename(columns={'GRADE': '告警级别'})
    
    # 转换时间为周并保留Period对象用于分组
    df['week_period'] = pd.to_datetime(df['alarm_first_time']).dt.to_period('W')
    
    # 创建自定义时间段显示格式
    def format_custom_period(period):
        start_date = period.start_time.strftime('%m.%d')
        end_date = (period.start_time + pd.Timedelta(days=6)).strftime('%m.%d')
        return f"{start_date}-{end_date}"
    
    # 处理自定义周期
    if custom_periods:
        # 解析自定义周期
        period_ranges = [p.strip() for p in custom_periods.split(',')]
        periods = []
        for pr in period_ranges:
            start, end = pr.split(' - ')
            # 转换为Period对象（确保为周粒度）
            start_date = pd.to_datetime(start).to_period('W')
            end_date = pd.to_datetime(end).to_period('W')
            # 生成连续的周范围
            current_period = start_date
            while current_period <= end_date:
                periods.append(current_period)
                current_period += 1  # 移动到下一周
        
        # 筛选数据
        df = df[df['week_period'].isin(periods)]
        
        # 如果没有数据匹配，提示用户
        if df.empty:
            print(f"警告: 没有数据匹配以下周期范围: {custom_periods}")
    
    # 创建字符串列用于显示(使用中文列名'周期')
    def format_period(period):
        if isinstance(period, str):
            # 处理week_period格式（如2025-06-30/2025-07-06）
            try:
                start, end = period.split('/')
                start_date = pd.to_datetime(start)
                end_date = pd.to_datetime(end)
                return f"{start_date.strftime('%m.%d')}-{end_date.strftime('%m.%d')}"
            except:
                return "00.00-00.00"
        else:
            # 处理Period对象
            start_date = period.start_time.strftime('%m.%d')
            end_date = (period.start_time + pd.Timedelta(days=6)).strftime('%m.%d')
            return f"{start_date}-{end_date}"
    
    # 统一处理周期列
    df['周期'] = df['week_period'].apply(format_period)
    
    if custom_periods:
        # 如果设置了自定义周期，则按设置的周期范围填充
        period_ranges = [p.strip() for p in custom_periods.split(',')]
        # 创建映射关系：week_period -> 对应的周期范围
        period_map = {}
        for pr in period_ranges:
            start, end = pr.split(' - ')
            start_date = pd.to_datetime(start).to_period('W')
            end_date = pd.to_datetime(end).to_period('W')
            current_period = start_date
            while current_period <= end_date:
                period_map[current_period] = pr
                current_period += 1
        # 填充周期列
        df['周期'] = df['week_period'].map(period_map)
    else:
        # 如果未设置自定义周期，则按默认周粒度填充
        df['周期'] = df['week_period'].apply(format_custom_period)
    
    # 确保告警数量列存在且不为空
    if 'ALARM_COUNT' not in df.columns:
        df['ALARM_COUNT'] = 1  # 如果不存在ALARM_COUNT列，则每条记录计为1
    df['ALARM_COUNT'] = df['ALARM_COUNT'].fillna(1).astype(int)
    
    # 确保处理状态列存在
    if 'deal_status' not in df.columns:
        df['deal_status'] = '未处理'  # 默认设为未处理
    
    # 添加调试输出
    print(f"准备写入文件到目录: {os.path.abspath(OUTPUT_DIR)}")
    # 支持输出前缀（多文件模式）
    output_name = f"{output_prefix}.xlsx" if output_prefix else "cleaned_data.xlsx"
    clean_path = os.path.join(OUTPUT_DIR, output_name)
    print(f"完整文件路径: {os.path.abspath(clean_path)}")
    
    try:
        # 确保输出目录存在
        print("正在创建输出目录(如果不存在)...")
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        print(f"目录已存在或创建成功: {os.path.abspath(OUTPUT_DIR)}")
        
        # 使用openpyxl直接创建Workbook
        from openpyxl import Workbook
        from openpyxl.utils.dataframe import dataframe_to_rows
        
        # 创建新的Workbook并删除默认sheet
        wb = Workbook()
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])
        
        # 将每个自定义周期的数据写入单独的sheet
        has_data = False
        if custom_periods:
            # 解析自定义周期
            period_ranges = [p.strip() for p in custom_periods.split(',')]
            for pr in period_ranges:
                start, end = pr.split(' - ')
                # 转换为Period对象
                start_date = pd.to_datetime(start).to_period('W')
                end_date = pd.to_datetime(end).to_period('W')
                # 筛选当前周期范围内的数据
                mask = (df['week_period'] >= start_date) & (df['week_period'] <= end_date)
                group = df[mask]
                if not group.empty:
                    sheet_name = f"每周告警统计_{start} - {end}"
                    ws = wb.create_sheet(title=sheet_name)
                    # 确保所有数据都是字符串或基本类型
                    group = group.astype(str)
                    # 写入数据
                    for r in dataframe_to_rows(group, index=False, header=True):
                        ws.append(r)
                    has_data = True
        else:
            # 默认按周粒度分组
            for week_period, group in df.groupby('week_period'):
                start_date = week_period.start_time.strftime('%Y-%m-%d')
                sheet_name = f"每周告警统计_{start_date}"
                ws = wb.create_sheet(title=sheet_name)
                group = group.astype(str)
                for r in dataframe_to_rows(group, index=False, header=True):
                    ws.append(r)
                has_data = True
        
        # 检查是否有数据
        if not has_data:
            print("警告: 没有数据匹配自定义周期范围，跳过文件保存")
            return None
        
        # 确保至少一个工作表可见
        if len(wb.sheetnames) == 0:
            wb.create_sheet(title="默认工作表")
        
        # 保存文件
        try:
            wb.save(clean_path)
        finally:
            wb.close()
        
        # 验证文件
        if not os.path.exists(clean_path):
            raise ValueError("文件未生成")
            
        file_size = os.path.getsize(clean_path)
        if file_size == 0:
            raise ValueError("生成的文件为空")
            
        print(f"成功生成清洗后的数据文件: {clean_path} (大小: {file_size}字节)")
        return clean_path
        
    except Exception as e:
        import traceback
        # 打印完整错误信息
        print("发生错误:")
        print(traceback.format_exc())
        
        # 如果出错，删除可能生成的不完整文件
        if os.path.exists(clean_path):
            try:
                print(f"尝试删除不完整文件: {clean_path}")
                os.remove(clean_path)
            except Exception as remove_error:
                print(f"删除文件失败: {str(remove_error)}")
        
        # 重新抛出错误确保可见
        print(f"写入Excel文件失败: {str(e)}")
        raise RuntimeError(f"写入Excel文件失败: {str(e)}") from e

if __name__ == "__main__":
    # 解析命令行参数
    parser = argparse.ArgumentParser(description='数据清洗脚本')
    parser.add_argument('--input-file', type=str, default=None,
                        help='输入Excel文件路径 (默认使用 config.INPUT_FILE)')
    parser.add_argument('--custom-periods', type=str, default=None,
                        help='自定义周期范围，格式为 mm.dd - mm.dd,mm.dd - mm.dd')
    parser.add_argument('--output-prefix', type=str, default=None,
                        help='输出文件名前缀（多文件模式使用）')
    args = parser.parse_args()
    
    print("=== 开始执行数据清洗脚本 ===")
    try:
        result_path = clean_data(custom_periods=args.custom_periods, input_file=args.input_file,
                                 output_prefix=args.output_prefix)
        print(f"=== 数据清洗完成，结果保存到: {result_path} ===")
    except Exception as e:
        print(f"!!! 脚本执行失败: {str(e)} !!!")
        raise